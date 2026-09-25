import csv
import re

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

GREEN = "217346"
GREEN_DARK = "1A5C38"
NUM = re.compile(r"^-?\d{1,3}(,\d{3})*(\.\d+)?$|^-?\d+(\.\d+)?$")


def _coerce(v):
    if not isinstance(v, str):
        return v
    s = v.strip()
    if not NUM.match(s) or re.match(r"^-?0\d", s):
        return v
    try:
        f = float(s.replace(",", ""))
    except ValueError:
        return v
    return int(f) if f == int(f) and "." not in s else f


def write_xlsx(table, path, title="Sheet1"):
    wb = Workbook()
    ws = wb.active
    ws.title = title[:31] or "Sheet1"
    cols = [c["name"] for c in table["columns"]]
    ws.append(cols)
    thin = Side(style="thin", color="D9D9D9")
    for i, c in enumerate(table["columns"], start=1):
        cell = ws.cell(row=1, column=i)
        cell.fill = PatternFill("solid", fgColor=GREEN_DARK if c["kind"] == "doc" else GREEN)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.alignment = Alignment(vertical="center")
        cell.border = Border(bottom=thin)
    for r in table["rows"]:
        ws.append([_coerce(r.get(c, "")) for c in cols])
    ws.freeze_panes = "A2"
    if cols:
        ws.auto_filter.ref = f"A1:{get_column_letter(len(cols))}{max(len(table['rows']) + 1, 1)}"
    for i, c in enumerate(cols, start=1):
        width = max([len(c)] + [len(str(r.get(c, ""))) for r in table["rows"][:200]])
        ws.column_dimensions[get_column_letter(i)].width = min(max(width + 2, 10), 50)
    wb.save(path)


def write_csv(table, path):
    cols = [c["name"] for c in table["columns"]]
    with open(path, "w", newline="", encoding="utf-8-sig") as f:
        w = csv.writer(f)
        w.writerow(cols)
        for r in table["rows"]:
            w.writerow([r.get(c, "") for c in cols])
